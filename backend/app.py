# backend/app.py
# Flask RESTful API server for ClashFree:
# Constraint-Based Timetable & Academic Resource Management System
# Supabase PostgreSQL + Fallback Storage, Multi-Department Resource Management,
# CSP+MRV Scheduler, Slot-Wise Occupancy, Resource Sharing & Workload Reports.

import os
import sys
import io
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from pathlib import Path
from dotenv import load_dotenv

root_dir = Path(__file__).resolve().parent.parent
load_dotenv(dotenv_path=root_dir / ".env")
load_dotenv()

backend_dir = os.path.dirname(os.path.abspath(__file__))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

import db
import excel_reader
import scheduler
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
CORS(app)

UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


# ---------------------------------------------------------------------------
# 1. HEALTH & SYSTEM STATUS
# ---------------------------------------------------------------------------
@app.route("/api/status", methods=["GET"])
def get_status():
    configured = db.is_configured()
    return jsonify({
        "status": "online",
        "system": "ClashFree Academic Resource & Timetable Management",
        "supabase_configured": configured,
        "message": "Connected to Supabase PostgreSQL" if configured else "Running in resilient local mode"
    })


# ---------------------------------------------------------------------------
# 2. DASHBOARD KPI SUMMARY
# ---------------------------------------------------------------------------
@app.route("/api/dashboard", methods=["GET"])
def get_dashboard():
    try:
        dept = request.args.get("department", "CSE")
        departments = db.get_all_departments()
        divisions = db.get_all_divisions(dept)
        batches = db.get_all_batches()
        faculty = db.get_all_faculty(dept)
        subjects = db.get_all_subjects(dept)
        rooms = db.get_all_rooms(dept)
        resource_cfg = db.get_resource_config(dept)
        timetable = db.get_timetable(department=dept)
        versions = db.get_timetable_versions()

        active_rooms, reserved_rooms, _ = scheduler.filter_and_reserve_resources(rooms, dept, resource_cfg)

        cr_total = len([r for r in rooms if r.get("type") == "Classroom"])
        cr_res = len([r for r in reserved_rooms if r.get("type") == "Classroom"])
        cr_active = cr_total - cr_res

        lab_total = len([r for r in rooms if r.get("type") == "Lab"])
        lab_res = len([r for r in reserved_rooms if r.get("type") == "Lab"])
        lab_active = lab_total - lab_res

        # Calculate currently occupied rooms across the timetable
        cr_occupied_set = set(r["room"] for r in timetable if not r.get("is_lab"))
        lab_occupied_set = set(r["room"] for r in timetable if r.get("is_lab"))

        cr_occupied = len([name for name in cr_occupied_set if any(r["name"] == name and r["type"] == "Classroom" for r in active_rooms)])
        cr_available = max(0, cr_active - cr_occupied)

        lab_occupied = len([name for name in lab_occupied_set if any(r["name"] == name and r["type"] == "Lab" for r in active_rooms)])
        lab_available = max(0, lab_active - lab_occupied)

        cr_occupancy_pct = round((cr_occupied / cr_active * 100), 1) if cr_active > 0 else 0
        lab_occupancy_pct = round((lab_occupied / lab_active * 100), 1) if lab_active > 0 else 0

        active_version = next((v for v in versions if v.get("is_active")), versions[0] if versions else None)

        return jsonify({
            "configured": db.is_configured(),
            "department": dept,
            "counts": {
                "departments": len(departments),
                "divisions": len(divisions),
                "batches": len(batches),
                "faculty": len(faculty),
                "subjects": len(subjects),
                "classrooms_total": cr_total,
                "classrooms_occupied": cr_occupied,
                "classrooms_available": cr_available,
                "classrooms_reserved": cr_res,
                "classrooms_occupancy_pct": cr_occupancy_pct,
                "labs_total": lab_total,
                "labs_occupied": lab_occupied,
                "labs_available": lab_available,
                "labs_reserved": lab_res,
                "labs_occupancy_pct": lab_occupancy_pct,
                "scheduled": len(timetable),
                "conflicts": 0,
                "active_version": active_version.get("name") if active_version else "Version 1"
            }
        })
    except Exception as e:
        return jsonify({"error": str(e), "configured": db.is_configured()}), 500


# ---------------------------------------------------------------------------
# 3. DEPARTMENTS CRUD
# ---------------------------------------------------------------------------
@app.route("/api/departments", methods=["GET", "POST"])
def manage_departments():
    if request.method == "GET":
        return jsonify(db.get_all_departments())
    elif request.method == "POST":
        payload = request.json or {}
        res = db.upsert_department(payload)
        return jsonify({"success": True, "data": res})


@app.route("/api/departments/<int:did>", methods=["DELETE"])
def delete_department_item(did):
    db.delete_department(did)
    return jsonify({"success": True})


# ---------------------------------------------------------------------------
# 4. ACADEMIC STRUCTURE
# ---------------------------------------------------------------------------
@app.route("/api/academic-structure", methods=["GET", "POST"])
def manage_academic_structure():
    dept = request.args.get("department", "CSE")
    if request.method == "GET":
        return jsonify(db.get_academic_structure(dept))
    elif request.method == "POST":
        res = db.upsert_academic_structure(request.json or {})
        return jsonify({"success": True, "data": res})


# ---------------------------------------------------------------------------
# 5. DIVISIONS & BATCHES CRUD
# ---------------------------------------------------------------------------
@app.route("/api/divisions", methods=["GET", "POST"])
def manage_divisions():
    dept = request.args.get("department")
    if request.method == "GET":
        return jsonify(db.get_all_divisions(dept))
    elif request.method == "POST":
        res = db.upsert_division(request.json or {})
        return jsonify({"success": True, "data": res})


@app.route("/api/divisions/<int:did>", methods=["DELETE"])
def delete_division_item(did):
    db.delete_division(did)
    return jsonify({"success": True})


@app.route("/api/batches", methods=["GET", "POST"])
def manage_batches():
    div_name = request.args.get("division")
    if request.method == "GET":
        return jsonify(db.get_all_batches(div_name))
    elif request.method == "POST":
        res = db.upsert_batch(request.json or {})
        return jsonify({"success": True, "data": res})


@app.route("/api/batches/<int:bid>", methods=["DELETE"])
def delete_batch_item(bid):
    db.delete_batch(bid)
    return jsonify({"success": True})


# ---------------------------------------------------------------------------
# 6. FACULTY CRUD
# ---------------------------------------------------------------------------
@app.route("/api/faculty", methods=["GET", "POST"])
def manage_faculty():
    dept = request.args.get("department")
    if request.method == "GET":
        return jsonify(db.get_all_faculty(dept))
    elif request.method == "POST":
        payload = request.json or {}
        if not payload.get("code"):
            # Auto generate initials short code (e.g. Vishal H. Pundikar -> VHP)
            name = payload.get("name", "").strip()
            parts = [p.strip(".") for p in name.split() if p.strip(".")]
            if len(parts) >= 3:
                payload["code"] = (parts[0][0] + parts[1][0] + parts[-1][0]).upper()
            elif len(parts) == 2:
                payload["code"] = (parts[0][0] + parts[1][0]).upper()
            elif len(parts) == 1 and len(parts[0]) >= 2:
                payload["code"] = parts[0][:2].upper()
            else:
                payload["code"] = "FAC"

        res = db.upsert_faculty(payload)
        return jsonify({"success": True, "data": res})


@app.route("/api/faculty/<int:fid>", methods=["DELETE"])
def delete_faculty_item(fid):
    db.delete_faculty(fid)
    return jsonify({"success": True})


# ---------------------------------------------------------------------------
# 7. ROOMS & RESOURCE RESERVATION
# ---------------------------------------------------------------------------
@app.route("/api/rooms", methods=["GET", "POST"])
def manage_rooms():
    dept = request.args.get("department")
    if request.method == "GET":
        return jsonify(db.get_all_rooms(dept))
    elif request.method == "POST":
        res = db.upsert_room(request.json or {})
        return jsonify({"success": True, "data": res})


@app.route("/api/rooms/<int:rid>", methods=["DELETE"])
def delete_room_item(rid):
    db.delete_room(rid)
    return jsonify({"success": True})


@app.route("/api/resource-config", methods=["GET", "POST"])
def manage_resource_config():
    dept = request.args.get("department", "CSE")
    if request.method == "GET":
        return jsonify(db.get_resource_config(dept))
    elif request.method == "POST":
        res = db.upsert_resource_config(request.json or {})
        return jsonify({"success": True, "data": res})


# ---------------------------------------------------------------------------
# 8. SUBJECTS CRUD
# ---------------------------------------------------------------------------
@app.route("/api/subjects", methods=["GET", "POST"])
def manage_subjects():
    dept = request.args.get("department")
    if request.method == "GET":
        return jsonify(db.get_all_subjects(dept))
    elif request.method == "POST":
        res = db.upsert_subject(request.json or {})
        return jsonify({"success": True, "data": res})


@app.route("/api/subjects/<int:sid>", methods=["DELETE"])
def delete_subject_item(sid):
    db.delete_subject(sid)
    return jsonify({"success": True})


# ---------------------------------------------------------------------------
# 9. TIME SLOTS
# ---------------------------------------------------------------------------
@app.route("/api/time-slots", methods=["GET", "POST"])
def manage_time_slots():
    if request.method == "GET":
        return jsonify(db.get_time_slots())
    elif request.method == "POST":
        db.save_time_slots(request.json or [])
        return jsonify({"success": True})


@app.route("/api/time-settings", methods=["GET", "POST"])
def manage_time_settings():
    if request.method == "GET":
        return jsonify(db.get_time_settings())
    elif request.method == "POST":
        res = db.save_time_settings(request.json or {})
        return jsonify({"success": True, "settings": res})


# ---------------------------------------------------------------------------
# 10. SLOT-WISE RESOURCE OCCUPANCY MODULE (HOD REQUIREMENT)
# ---------------------------------------------------------------------------
@app.route("/api/occupancy", methods=["GET"])
def get_slot_occupancy():
    """
    Returns slot-wise resource status: OCCUPIED, AVAILABLE, RESERVED.
    Query params: day, period, department, type
    """
    day = request.args.get("day", "Monday")
    period = request.args.get("period", "09:00-10:00")
    dept = request.args.get("department", "CSE")
    r_type = request.args.get("type", "ALL")

    all_rooms = db.get_all_rooms(dept)
    resource_cfg = db.get_resource_config(dept)
    timetable = db.get_timetable(department=dept)
    sharing_reqs = db.get_resource_sharing_requests()

    active_rooms, reserved_rooms, _ = scheduler.filter_and_reserve_resources(all_rooms, dept, resource_cfg)
    reserved_room_names = set(r["name"] for r in reserved_rooms)

    # Filter by slot in timetable
    slot_sessions = [r for r in timetable if r.get("day") == day and r.get("period") == period]
    occupied_map = {r["room"]: r for r in slot_sessions}

    # Filter sharing requests approved for this day/period
    approved_sharing = {
        r["resource_name"]: r for r in sharing_reqs
        if r.get("day") == day and r.get("period") == period and r.get("status") == "APPROVED"
    }

    results = []
    cr_occ, cr_free, cr_res = 0, 0, 0
    lab_occ, lab_free, lab_res = 0, 0, 0

    for room in all_rooms:
        if r_type != "ALL" and room.get("type") != r_type:
            continue

        r_name = room.get("name")
        is_lab = (room.get("type") == "Lab")

        if r_name in reserved_room_names or room.get("is_reserved"):
            status = "RESERVED"
            allocation = {
                "division": "-",
                "batch": "-",
                "subject": "-",
                "faculty": "-",
                "notes": room.get("reservation_reason") or "HOD Reserved Resource (Exempt from normal scheduling)"
            }
            if is_lab:
                lab_res += 1
            else:
                cr_res += 1

        elif r_name in occupied_map:
            sess = occupied_map[r_name]
            status = "OCCUPIED"
            allocation = {
                "division": sess.get("section"),
                "batch": sess.get("batch") or "Whole Class",
                "subject": sess.get("subject"),
                "faculty": f"{sess.get('faculty_name', '')} ({sess.get('faculty_code', '')})",
                "notes": "Regular Timetable Session"
            }
            if is_lab:
                lab_occ += 1
            else:
                cr_occ += 1

        elif r_name in approved_sharing:
            sh = approved_sharing[r_name]
            status = "OCCUPIED"
            allocation = {
                "division": sh.get("requesting_department"),
                "batch": "Shared Allocation",
                "subject": sh.get("reason"),
                "faculty": sh.get("requested_by"),
                "notes": f"Cross-Department Shared with {sh.get('requesting_department')}"
            }
            if is_lab:
                lab_occ += 1
            else:
                cr_occ += 1

        else:
            status = "AVAILABLE"
            allocation = {
                "division": "-",
                "batch": "-",
                "subject": "-",
                "faculty": "-",
                "notes": "Free for Department / Available for Cross-Dept Sharing"
            }
            if is_lab:
                lab_free += 1
            else:
                cr_free += 1

        results.append({
            "resource": r_name,
            "type": room.get("type"),
            "department": room.get("department", "CSE"),
            "capacity": room.get("capacity"),
            "num_pcs": room.get("num_pcs", 0),
            "status": status,
            "allocation": allocation
        })

    cr_total = cr_occ + cr_free + cr_res
    lab_total = lab_occ + lab_free + lab_res

    summary = {
        "classrooms": {
            "total": cr_total,
            "occupied": cr_occ,
            "available": cr_free,
            "reserved": cr_res,
            "occupancy_pct": round((cr_occ / max(1, cr_total - cr_res) * 100), 1) if (cr_total - cr_res) > 0 else 0
        },
        "labs": {
            "total": lab_total,
            "occupied": lab_occ,
            "available": lab_free,
            "reserved": lab_res,
            "occupancy_pct": round((lab_occ / max(1, lab_total - lab_res) * 100), 1) if (lab_total - lab_res) > 0 else 0
        }
    }

    return jsonify({
        "day": day,
        "period": period,
        "department": dept,
        "summary": summary,
        "resources": results
    })


# ---------------------------------------------------------------------------
# 11. RESOURCE SHARING BETWEEN DEPARTMENTS
# ---------------------------------------------------------------------------
@app.route("/api/resource-sharing", methods=["GET", "POST", "PUT"])
def manage_resource_sharing():
    if request.method == "GET":
        return jsonify(db.get_resource_sharing_requests())

    elif request.method == "POST":
        payload = request.json or {}
        res_name = payload.get("resource_name")
        day = payload.get("day")
        period = payload.get("period")
        req_dept = payload.get("requesting_department")

        # Validate that the resource is NOT already occupied in timetable or reserved
        timetable = db.get_timetable()
        rooms = db.get_all_rooms()
        target_room = next((r for r in rooms if r["name"] == res_name), None)

        if not target_room:
            return jsonify({"error": "Resource not found"}), 404

        if target_room.get("is_reserved"):
            return jsonify({"error": f"Resource {res_name} is reserved and cannot be requested"}), 400

        # Check existing timetable booking
        clash = any(t for t in timetable if t.get("day") == day and t.get("period") == period and t.get("room") == res_name)
        if clash:
            return jsonify({"error": f"Resource {res_name} is occupied in the master timetable during {day} {period}"}), 400

        # Check other approved sharing
        existing_reqs = db.get_resource_sharing_requests()
        for r in existing_reqs:
            if r.get("resource_name") == res_name and r.get("day") == day and r.get("period") == period and r.get("status") == "APPROVED":
                return jsonify({"error": f"Resource {res_name} is already approved for {r.get('requesting_department')} during this slot"}), 400

        payload["owner_department"] = target_room.get("department", "CSE")
        new_req = db.create_resource_sharing_request(payload)
        return jsonify({"success": True, "request": new_req})

    elif request.method == "PUT":
        payload = request.json or {}
        req_id = payload.get("id")
        status = payload.get("status")  # APPROVED or REJECTED
        reviewed_by = payload.get("reviewed_by", "HOD")
        db.update_resource_sharing_request(req_id, status, reviewed_by)
        return jsonify({"success": True})


# ---------------------------------------------------------------------------
# 12. TIMETABLE & SCHEDULER (CSP + MRV + BACKTRACKING)
# ---------------------------------------------------------------------------
@app.route("/api/timetable", methods=["GET"])
def get_timetable_data():
    dept = request.args.get("department", "CSE")
    version_id = request.args.get("version_id")
    data = db.get_timetable(version_id=version_id, department=dept)
    return jsonify(data)


@app.route("/api/generate-timetable", methods=["POST"])
def run_scheduler():
    try:
        payload = request.json or {}
        dept = payload.get("department", "CSE")
        version_name = payload.get("version_name") or f"Version {len(db.get_timetable_versions()) + 1}"

        faculty = db.get_all_faculty(dept)
        rooms = db.get_all_rooms(dept)
        subjects = db.get_all_subjects(dept)
        divisions = db.get_all_divisions(dept)
        batches = db.get_all_batches()
        time_slots = db.get_time_slots()
        resource_cfg = db.get_resource_config(dept)
        sharing_reqs = db.get_resource_sharing_requests()
        time_settings = db.get_time_settings()
        days = payload.get("days") or time_settings.get("working_days", ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])

        success, message, timetable_rows, conflicts, meta = scheduler.generate_timetable(
            faculty_list=faculty,
            room_list=rooms,
            subject_list=subjects,
            division_list=divisions,
            batch_list=batches,
            time_slots=time_slots,
            resource_config=resource_cfg,
            department=dept,
            approved_sharing=sharing_reqs,
            days=days
        )

        if success:
            # Create new timetable version
            new_v_id = len(db.get_timetable_versions()) + 1
            new_version = {
                "id": new_v_id,
                "version_number": new_v_id,
                "name": version_name,
                "academic_year": "2026-2027",
                "semester": "Semester III",
                "department": dept,
                "description": f"Generated with {meta.get('reserved_rooms_count')} reserved rooms and {len(timetable_rows)} sessions",
                "is_active": True,
                "created_by": payload.get("created_by", "HOD")
            }
            db.create_timetable_version(new_version)
            db.set_active_version(new_v_id)
            db.save_timetable(timetable_rows, version_id=new_v_id, department=dept)

            return jsonify({
                "success": True,
                "message": message,
                "count": len(timetable_rows),
                "timetable": timetable_rows,
                "version": new_version,
                "metadata": meta,
                "conflicts": []
            })
        else:
            return jsonify({
                "success": False,
                "message": message,
                "conflicts": conflicts,
                "metadata": meta
            }), 422

    except Exception as e:
        return jsonify({"success": False, "message": str(e), "conflicts": [str(e)]}), 500


# ---------------------------------------------------------------------------
# 13. TIMETABLE VERSIONS
# ---------------------------------------------------------------------------
@app.route("/api/timetable-versions", methods=["GET", "POST"])
def manage_versions():
    if request.method == "GET":
        return jsonify(db.get_timetable_versions())
    elif request.method == "POST":
        payload = request.json or {}
        vid = payload.get("version_id")
        if vid:
            db.set_active_version(int(vid))
            return jsonify({"success": True, "active_version_id": vid})
        return jsonify({"error": "Version ID required"}), 400


# ---------------------------------------------------------------------------
# 14. FACULTY WORKLOAD & ACADEMIC LOAD REPORT (HOD REQUIREMENT)
# ---------------------------------------------------------------------------
@app.route("/api/faculty-workload", methods=["GET"])
def get_faculty_workload():
    dept = request.args.get("department", "CSE")
    faculty_list = db.get_all_faculty(dept)
    timetable = db.get_timetable(department=dept)

    workload_report = []

    for fac in faculty_list:
        code = fac.get("code")
        name = fac.get("name")
        designation = fac.get("designation", "Assistant Professor")
        max_load = fac.get("max_load", 18)

        fac_sessions = [r for r in timetable if r.get("faculty_code") == code]

        lec_sessions = [r for r in fac_sessions if not r.get("is_lab")]
        lab_sessions = [r for r in fac_sessions if r.get("is_lab")]

        lec_hours = len(lec_sessions)
        lab_hours = len(lab_sessions) * 1  # 1 hour per session
        total_load = lec_hours + lab_hours

        # Group by subject and division/batch
        breakdown = {}
        for s in fac_sessions:
            subj = s.get("subject")
            div_batch = s.get("section") + (f" ({s.get('batch')})" if s.get("batch") else "")
            key = f"{subj} - {div_batch}"
            breakdown[key] = breakdown.get(key, 0) + 1

        status = "Optimal"
        if total_load > max_load:
            status = "Overloaded"
        elif total_load < max(0, max_load - 6):
            status = "Underloaded"

        workload_report.append({
            "faculty_name": name,
            "short_code": code,
            "designation": designation,
            "department": fac.get("department", dept),
            "lecture_hours": lec_hours,
            "lab_hours": lab_hours,
            "total_hours": total_load,
            "max_load": max_load,
            "status": status,
            "breakdown": breakdown
        })

    return jsonify(workload_report)


@app.route("/api/academic-load-report", methods=["GET"])
def get_academic_load_report():
    """
    Format inspired by college reference table:
    Faculty Name | Designation | Class/Year | Div/Batch | Course | Lec Hours | Lab Hours | Direct Contact | Project | Seminar | Grand Total
    """
    dept = request.args.get("department", "CSE")
    faculty_list = db.get_all_faculty(dept)
    timetable = db.get_timetable(department=dept)

    rows = []
    for fac in faculty_list:
        code = fac.get("code")
        fac_sessions = [r for r in timetable if r.get("faculty_code") == code]

        # Group by course and division/batch
        grouped = {}
        for s in fac_sessions:
            key = (s.get("subject"), s.get("section"), s.get("batch") or "Whole Class")
            if key not in grouped:
                grouped[key] = {"lec": 0, "lab": 0}
            if s.get("is_lab"):
                grouped[key]["lab"] += 1
            else:
                grouped[key]["lec"] += 1

        if not grouped:
            # Faculty with 0 teaching sessions
            rows.append({
                "faculty_name": fac.get("name"),
                "short_code": code,
                "designation": fac.get("designation", "Assistant Professor"),
                "class_year": "Second Year",
                "division_batch": "-",
                "course_name": "-",
                "lecture_hours": 0,
                "lab_hours": 0,
                "direct_contact_hours": 0,
                "project_hours": 2,
                "seminar_hours": 1,
                "internship_hours": 0,
                "grand_total": 3
            })
        else:
            for (course, div, batch), counts in grouped.items():
                lec_h = counts["lec"]
                lab_h = counts["lab"]
                direct_h = lec_h + lab_h
                rows.append({
                    "faculty_name": fac.get("name"),
                    "short_code": code,
                    "designation": fac.get("designation", "Assistant Professor"),
                    "class_year": "Second Year (Sem III)",
                    "division_batch": f"{div} / {batch}",
                    "course_name": course,
                    "lecture_hours": lec_h,
                    "lab_hours": lab_h,
                    "direct_contact_hours": direct_h,
                    "project_hours": 2 if "Professor" in fac.get("designation", "") else 0,
                    "seminar_hours": 1,
                    "internship_hours": 0,
                    "grand_total": direct_h + (2 if "Professor" in fac.get("designation", "") else 0) + 1
                })

    return jsonify(rows)


# ---------------------------------------------------------------------------
# 15. CONFLICT DIAGNOSTICS
# ---------------------------------------------------------------------------
@app.route("/api/conflicts", methods=["GET"])
def get_conflicts():
    try:
        dept = request.args.get("department", "CSE")
        faculty = db.get_all_faculty(dept)
        rooms = db.get_all_rooms(dept)
        subjects = db.get_all_subjects(dept)
        divisions = db.get_all_divisions(dept)
        batches = db.get_all_batches()
        time_slots = db.get_time_slots()
        resource_cfg = db.get_resource_config(dept)

        success, _, _, conflicts, meta = scheduler.generate_timetable(
            faculty, rooms, subjects, divisions, batches, time_slots, resource_cfg, dept
        )

        active_rooms, reserved_rooms, _ = scheduler.filter_and_reserve_resources(rooms, dept, resource_cfg)

        issues = list(conflicts)
        # Check capacity warnings
        for div in divisions:
            cap_fit = [r for r in active_rooms if r.get("type") == "Classroom" and r.get("capacity", 0) >= div.get("strength", 60)]
            if not cap_fit:
                issues.append(f"CAPACITY ISSUE: Division '{div.get('name')}' strength ({div.get('strength')}) exceeds all available non-reserved classrooms.")

        # Check PC ratio warnings
        for b in batches:
            pc_fit = [r for r in active_rooms if r.get("type") == "Lab" and r.get("num_pcs", 0) >= b.get("strength", 30)]
            if not pc_fit:
                issues.append(f"PC CAPACITY ISSUE: Lab Batch '{b.get('name')}' strength ({b.get('strength')}) exceeds PC count in available non-reserved labs.")

        return jsonify({
            "conflicts": issues,
            "has_conflicts": len(issues) > 0,
            "active_rooms_count": len(active_rooms),
            "reserved_rooms_count": len(reserved_rooms)
        })
    except Exception as e:
        return jsonify({"conflicts": [f"Error checking conflicts: {str(e)}"], "has_conflicts": True})


# ---------------------------------------------------------------------------
# 16. EXCEL IMPORT
# ---------------------------------------------------------------------------
@app.route("/api/import-excel", methods=["POST"])
def import_excel():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    file = request.files["file"]
    if not file.filename.endswith(".xlsx"):
        return jsonify({"error": "Only .xlsx files are supported"}), 400

    file_bytes = io.BytesIO(file.read())
    data_dict, errors, warnings, summary = excel_reader.read_excel(file_bytes)
    if errors:
        return jsonify({"errors": errors, "warnings": warnings, "summary": summary}), 400

    try:
        if data_dict.get("faculty"):
            db.insert_faculty(data_dict["faculty"])
        if data_dict.get("rooms"):
            db.insert_rooms(data_dict["rooms"])
        if data_dict.get("subjects"):
            db.insert_subjects(data_dict["subjects"])
        if data_dict.get("sections"):
            db.insert_sections(data_dict["sections"])
        if data_dict.get("divisions"):
            db.insert_divisions(data_dict["divisions"])
        if data_dict.get("batches"):
            db.insert_batches(data_dict["batches"])
        if data_dict.get("departments"):
            for dep in data_dict["departments"]:
                db.upsert_department(dep)
    except Exception as e:
        print(f"Import DB warning: {e}")

    return jsonify({
        "success": True,
        "summary": summary,
        "warnings": warnings
    })


# ---------------------------------------------------------------------------
# 17. EXCEL EXPORT (TIMETABLE & ACADEMIC LOAD)
# ---------------------------------------------------------------------------
@app.route("/api/export/timetable", methods=["GET"])
def export_timetable_excel():
    dept = request.args.get("department", "CSE")
    timetable = db.get_timetable(department=dept)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{dept} Master Timetable"

    headers = ["Day", "Time Slot", "Department", "Division", "Batch", "Subject", "Faculty Code", "Faculty Name", "Room", "Type"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in timetable:
        ws.append([
            row.get("day"),
            row.get("period"),
            row.get("department", dept),
            row.get("section"),
            row.get("batch") or "Whole Class",
            row.get("subject"),
            row.get("faculty_code"),
            row.get("faculty_name"),
            row.get("room"),
            "LAB" if row.get("is_lab") else "LECTURE"
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{dept}_Master_Timetable.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/api/export/workload", methods=["GET"])
def export_workload_excel():
    dept = request.args.get("department", "CSE")
    res = get_academic_load_report()
    data = res.get_json()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Academic Load Report"

    headers = ["Faculty Name", "Code", "Designation", "Class/Year", "Div/Batch", "Course", "Lec Hrs", "Lab Hrs", "Direct Contact", "Project", "Seminar", "Grand Total"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in data:
        ws.append([
            r.get("faculty_name"),
            r.get("short_code"),
            r.get("designation"),
            r.get("class_year"),
            r.get("division_batch"),
            r.get("course_name"),
            r.get("lecture_hours"),
            r.get("lab_hours"),
            r.get("direct_contact_hours"),
            r.get("project_hours"),
            r.get("seminar_hours"),
            r.get("grand_total")
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{dept}_Academic_Load_Report.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    port = int(os.getenv("PORT", 5000))
    print(f"ClashFree Backend API running on http://127.0.0.1:{port}")
    app.run(host="0.0.0.0", port=port, debug=True)
