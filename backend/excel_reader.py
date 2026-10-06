# backend/excel_reader.py
# Reads uploaded Excel file and returns validated data.
# Preserved from original with improvements for Supabase-compatible format.

import openpyxl


def read_excel(filepath: str):
    """
    Read .xlsx file and return (data_dict, errors_list).
    data_dict keys: faculty, rooms, subjects, sections
    errors_list: plain English error messages for the user.
    """
    errors = []

    try:
        workbook = openpyxl.load_workbook(filepath)
    except Exception as e:
        return None, [f"Could not open Excel file: {str(e)}"]

    # ---- Faculty sheet ----
    faculty_data = []
    if "Faculty" not in workbook.sheetnames:
        errors.append("Sheet 'Faculty' is missing from the Excel file.")
    else:
        sheet = workbook["Faculty"]
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue
            name           = str(row[0]).strip() if len(row) > 0 and row[0] else ""
            subjects_raw   = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            available_from = str(row[2]).strip() if len(row) > 2 and row[2] else "09:00"
            available_to   = str(row[3]).strip() if len(row) > 3 and row[3] else "17:00"
            preferred_time = str(row[4]).strip() if len(row) > 4 and row[4] else "Any"
            max_load       = row[5] if len(row) > 5 and row[5] is not None else 20

            if not name or name == "None":
                errors.append(f"Faculty sheet, Row {row_num}: Faculty name is missing.")
                continue

            subjects = [s.strip() for s in subjects_raw.split(",") if s.strip()]
            preferred_time = preferred_time.capitalize()
            if preferred_time not in ("Morning", "Afternoon", "Any"):
                preferred_time = "Any"

            try:
                max_load = int(max_load)
            except (ValueError, TypeError):
                max_load = 20

            faculty_data.append({
                "name":           name,
                "subjects":       subjects,
                "available_from": available_from,
                "available_to":   available_to,
                "preferred_time": preferred_time,
                "max_load":       max_load,
                "code":           "",  # filled by generate_faculty_codes()
            })

    # ---- Rooms sheet ----
    room_data = []
    if "Rooms" not in workbook.sheetnames:
        errors.append("Sheet 'Rooms' is missing from the Excel file.")
    else:
        sheet = workbook["Rooms"]
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue
            room_name = str(row[0]).strip() if row[0] else ""
            room_type = str(row[1]).strip() if row[1] else ""
            capacity  = row[2] if row[2] else 0

            if not room_name or room_name == "None":
                errors.append(f"Rooms sheet, Row {row_num}: Room name is missing.")
                continue
            if room_type not in ("Classroom", "Lab"):
                errors.append(
                    f"Rooms sheet, Row {row_num}: Type must be 'Classroom' or 'Lab'. Got '{room_type}'."
                )
                continue
            try:
                capacity = int(capacity)
            except (ValueError, TypeError):
                errors.append(f"Rooms sheet, Row {row_num}: Capacity must be a number.")
                continue

            room_data.append({"name": room_name, "type": room_type, "capacity": capacity})

    # ---- Subjects sheet ----
    subject_data = []
    if "Subjects" not in workbook.sheetnames:
        errors.append("Sheet 'Subjects' is missing from the Excel file.")
    else:
        sheet = workbook["Subjects"]
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue
            subject_name = str(row[0]).strip() if row[0] else ""
            lec_pw       = row[1] if row[1] is not None else 0
            lab_pw       = row[2] if row[2] is not None else 0
            req_lab      = str(row[3]).strip().lower() if row[3] else "no"

            if not subject_name or subject_name == "None":
                errors.append(f"Subjects sheet, Row {row_num}: Subject name is missing.")
                continue
            try:
                lec_pw = int(lec_pw)
                lab_pw = int(lab_pw)
            except (ValueError, TypeError):
                errors.append(f"Subjects sheet, Row {row_num}: Lectures/Labs per week must be numbers.")
                continue

            subject_data.append({
                "name":             subject_name,
                "lectures_per_week": lec_pw,
                "labs_per_week":    lab_pw,
                "requires_lab":     req_lab in ("yes", "true", "1"),
            })

    # ---- Sections sheet ----
    section_data = []
    if "Sections" not in workbook.sheetnames:
        errors.append("Sheet 'Sections' is missing from the Excel file.")
    else:
        sheet = workbook["Sections"]
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue
            section_name = str(row[0]).strip() if row[0] else ""
            strength     = row[1] if row[1] is not None else 0

            if not section_name or section_name == "None":
                errors.append(f"Sections sheet, Row {row_num}: Section name is missing.")
                continue
            try:
                strength = int(strength)
            except (ValueError, TypeError):
                errors.append(f"Sections sheet, Row {row_num}: Strength must be a number.")
                continue

            section_data.append({"name": section_name, "strength": strength})

    return {
        "faculty":  faculty_data,
        "rooms":    room_data,
        "subjects": subject_data,
        "sections": section_data,
    }, errors


def generate_faculty_codes(faculty_data: list) -> list:
    """
    Auto-generate unique short codes.
    Kunal Patil → KP
    If duplicate → KP1, KP2, ...
    """
    code_count = {}
    for f in faculty_data:
        parts = f["name"].strip().split()
        base  = (parts[0][0] + parts[-1][0]).upper() if len(parts) >= 2 else parts[0][:2].upper()
        code_count[base] = code_count.get(base, 0) + 1

    occurrence = {}
    for f in faculty_data:
        parts = f["name"].strip().split()
        base  = (parts[0][0] + parts[-1][0]).upper() if len(parts) >= 2 else parts[0][:2].upper()
        if code_count[base] == 1:
            f["code"] = base
        else:
            occurrence[base] = occurrence.get(base, 0) + 1
            f["code"] = base + str(occurrence[base])

    return faculty_data

