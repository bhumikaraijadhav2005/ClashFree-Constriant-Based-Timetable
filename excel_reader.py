# excel_reader.py
# This file reads the uploaded Excel file and converts it into Python data.
# It also validates the data and returns easy-to-understand error messages.

import openpyxl


def read_excel(filepath):
    """
    Read the Excel file at 'filepath'.
    Returns a tuple: (data_dict, errors_list)

    data_dict contains:
        "faculty"  : list of faculty dicts
        "rooms"    : list of room dicts
        "subjects" : list of subject dicts
        "sections" : list of section dicts

    errors_list is a list of plain English error messages.
    If errors_list is not empty, the data should NOT be used.
    """

    errors = []

    try:
        workbook = openpyxl.load_workbook(filepath)
    except Exception as e:
        return None, [f"Could not open Excel file: {str(e)}"]

    # -----------------------------------------------------------------------
    # Read Faculty sheet
    # -----------------------------------------------------------------------
    faculty_data = []
    if "Faculty" not in workbook.sheetnames:
        errors.append("Sheet named 'Faculty' is missing from the Excel file.")
    else:
        sheet = workbook["Faculty"]
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            # Skip completely empty rows
            if all(cell is None for cell in row):
                continue

            name          = str(row[0]).strip() if row[0] else ""
            subjects_raw  = str(row[1]).strip() if row[1] else ""
            available_from = str(row[2]).strip() if row[2] else "09:00"
            available_to   = str(row[3]).strip() if row[3] else "17:00"
            preferred_time = str(row[4]).strip() if row[4] else "Any"

            if not name or name == "None":
                errors.append(f"Faculty sheet, Row {row_num}: Faculty name is missing.")
                continue

            # Subjects can be comma-separated: "DSA, OS, CN"
            subjects = [s.strip() for s in subjects_raw.split(",") if s.strip()]

            # Normalize preferred_time
            preferred_time = preferred_time.capitalize()
            if preferred_time not in ("Morning", "Afternoon", "Any"):
                preferred_time = "Any"

            faculty_data.append({
                "name": name,
                "subjects": subjects,
                "available_from": available_from,
                "available_to": available_to,
                "preferred_time": preferred_time,
                "code": ""  # will be filled by data.generate_faculty_codes()
            })

    # -----------------------------------------------------------------------
    # Read Rooms sheet
    # -----------------------------------------------------------------------
    room_data = []
    if "Rooms" not in workbook.sheetnames:
        errors.append("Sheet named 'Rooms' is missing from the Excel file.")
    else:
        sheet = workbook["Rooms"]
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue

            room_name = str(row[0]).strip() if row[0] else ""
            room_type = str(row[1]).strip() if row[1] else ""
            capacity  = row[2] if row[2] else 0

            if not room_name or room_name == "None":
                errors.append(f"Rooms sheet, Row {row_num}: Room name is missing.")
                continue

            if room_type not in ("Classroom", "Lab"):
                errors.append(
                    f"Rooms sheet, Row {row_num}: Room type must be 'Classroom' or 'Lab'. Got '{room_type}'."
                )
                continue

            try:
                capacity = int(capacity)
            except (ValueError, TypeError):
                errors.append(f"Rooms sheet, Row {row_num}: Capacity must be a number.")
                continue

            room_data.append({
                "name": room_name,
                "type": room_type,
                "capacity": capacity
            })

    # -----------------------------------------------------------------------
    # Read Subjects sheet
    # -----------------------------------------------------------------------
    subject_data = []
    if "Subjects" not in workbook.sheetnames:
        errors.append("Sheet named 'Subjects' is missing from the Excel file.")
    else:
        sheet = workbook["Subjects"]
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue

            subject_name    = str(row[0]).strip() if row[0] else ""
            lectures_per_wk = row[1] if row[1] is not None else 0
            labs_per_wk     = row[2] if row[2] is not None else 0
            requires_lab    = str(row[3]).strip().lower() if row[3] else "no"

            if not subject_name or subject_name == "None":
                errors.append(f"Subjects sheet, Row {row_num}: Subject name is missing.")
                continue

            try:
                lectures_per_wk = int(lectures_per_wk)
                labs_per_wk     = int(labs_per_wk)
            except (ValueError, TypeError):
                errors.append(
                    f"Subjects sheet, Row {row_num}: LecturesPerWeek and LabsPerWeek must be numbers."
                )
                continue

            subject_data.append({
                "name": subject_name,
                "lectures_per_week": lectures_per_wk,
                "labs_per_week": labs_per_wk,
                "requires_lab": requires_lab in ("yes", "true", "1")
            })

    # -----------------------------------------------------------------------
    # Read Sections sheet
    # -----------------------------------------------------------------------
    section_data = []
    if "Sections" not in workbook.sheetnames:
        errors.append("Sheet named 'Sections' is missing from the Excel file.")
    else:
        sheet = workbook["Sections"]
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue

            section_name = str(row[0]).strip() if row[0] else ""
            strength     = row[1] if row[1] is not None else 0

            if not section_name or section_name == "None":
                errors.append(f"Sections sheet, Row {row_num}: Section name is missing.")
                continue

            try:
                strength = int(strength)
            except (ValueError, TypeError):
                errors.append(f"Sections sheet, Row {row_num}: Strength must be a number.")
                continue

            section_data.append({
                "name": section_name,
                "strength": strength
            })

    # Return all data
    all_data = {
        "faculty":  faculty_data,
        "rooms":    room_data,
        "subjects": subject_data,
        "sections": section_data
    }
    return all_data, errors
