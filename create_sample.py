# create_sample.py
# Run this once to create sample_data.xlsx for testing.
# Command: python create_sample.py

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()

# ----------------------------------------------------------------
# Helper to style header rows
# ----------------------------------------------------------------
def style_header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1a73e8")
        cell.alignment = Alignment(horizontal="center")

# ----------------------------------------------------------------
# Sheet 1: Faculty
# Columns: Name | Subjects | AvailableFrom | AvailableTo | PreferredTime
# ----------------------------------------------------------------
ws_faculty = wb.active
ws_faculty.title = "Faculty"

style_header(ws_faculty, ["Name", "Subjects", "AvailableFrom", "AvailableTo", "PreferredTime"])

faculty_data = [
    # Kunal Patil and Kiran Patil both get "KP" → will become KP1 and KP2
    ["Kunal Patil",      "DSA, OS",      "09:00", "16:00", "Morning"],
    ["Kiran Patil",      "DBMS, CN",     "09:00", "17:00", "Morning"],
    ["Amit Sharma",      "Maths, TOC",   "09:00", "17:00", "Any"],
    ["Sneha Desai",      "DSA, CN",      "11:00", "17:00", "Afternoon"],
    ["Rahul Mehta",      "OS, TOC",      "09:00", "17:00", "Any"],
    ["Priya Joshi",      "DBMS, Maths",  "09:00", "14:00", "Morning"],
]

for row in faculty_data:
    ws_faculty.append(row)

# ----------------------------------------------------------------
# Sheet 2: Rooms
# Columns: RoomName | Type | Capacity
# ----------------------------------------------------------------
ws_rooms = wb.create_sheet("Rooms")

style_header(ws_rooms, ["RoomName", "Type", "Capacity"])

rooms_data = [
    ["Room101",   "Classroom", 65],
    ["Room102",   "Classroom", 65],
    ["Room103",   "Classroom", 60],
    ["Room104",   "Classroom", 60],
    ["Lab1",      "Lab",       65],
    ["Lab2",      "Lab",       65],
]

for row in rooms_data:
    ws_rooms.append(row)

# ----------------------------------------------------------------
# Sheet 3: Subjects
# Columns: Subject | LecturesPerWeek | LabsPerWeek | RequiresLab
# ----------------------------------------------------------------
ws_subjects = wb.create_sheet("Subjects")

style_header(ws_subjects, ["Subject", "LecturesPerWeek", "LabsPerWeek", "RequiresLab"])

subjects_data = [
    ["DSA",   3, 0, "No"],
    ["DBMS",  2, 1, "Yes"],
    ["OS",    2, 0, "No"],
    ["CN",    2, 0, "No"],
    ["Maths", 2, 0, "No"],
    ["TOC",   2, 0, "No"],
]

for row in subjects_data:
    ws_subjects.append(row)

# ----------------------------------------------------------------
# Sheet 4: Sections
# Columns: SectionName | Strength
# ----------------------------------------------------------------
ws_sections = wb.create_sheet("Sections")

style_header(ws_sections, ["SectionName", "Strength"])

sections_data = [
    ["CSE-A", 60],
    ["CSE-B", 58],
]

for row in sections_data:
    ws_sections.append(row)

# ----------------------------------------------------------------
# Set column widths for readability
# ----------------------------------------------------------------
for ws in [ws_faculty, ws_rooms, ws_subjects, ws_sections]:
    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

# ----------------------------------------------------------------
# Save
# ----------------------------------------------------------------
wb.save("sample_data.xlsx")
print("sample_data.xlsx created successfully!")
print()
print("Data summary:")
print(f"  Faculty:  {len(faculty_data)}")
print(f"  Rooms:    {len(rooms_data)}")
print(f"  Subjects: {len(subjects_data)}")
print(f"  Sections: {len(sections_data)}")
print()
print("Note: 'Kunal Patil' and 'Kiran Patil' both start with KP.")
print("      The system will automatically assign them KP1 and KP2.")
